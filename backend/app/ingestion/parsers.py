from dataclasses import dataclass
from pathlib import Path

import fitz
from pptx import Presentation

from ..core.paths import pages_dir


@dataclass(frozen=True)
class ParsedPage:
    page_number: int
    text: str
    image_path: str | None = None


def parse_pdf(path: Path, document_id: int) -> list[ParsedPage]:
    parsed_pages: list[ParsedPage] = []
    output_dir = pages_dir() / str(document_id)
    output_dir.mkdir(parents=True, exist_ok=True)

    with fitz.open(path) as pdf:
        for index, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            image_path = output_dir / f"page-{index}.png"
            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            pixmap.save(image_path)
            parsed_pages.append(
                ParsedPage(page_number=index, text=text, image_path=str(image_path))
            )
    return parsed_pages


def parse_pptx(path: Path) -> list[ParsedPage]:
    presentation = Presentation(path)
    parsed_pages: list[ParsedPage] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text.strip())
        parsed_pages.append(ParsedPage(page_number=index, text="\n".join(texts).strip()))
    return parsed_pages


def parse_document(path: str, file_type: str, document_id: int) -> list[ParsedPage]:
    file_path = Path(path)
    if file_type == "pdf":
        return parse_pdf(file_path, document_id)
    if file_type == "pptx":
        return parse_pptx(file_path)
    raise ValueError(f"Unsupported file type: {file_type}")
